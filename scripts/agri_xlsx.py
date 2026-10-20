"""Read ELSTAT spreadsheets (.xlsx and the older .xls) into plain grids, with merged cells.

ELSTAT names both formats ".xls", so the format is detected from the file content, not the extension.
.xlsx files are read with openpyxl and .xls files with xlrd (pip install openpyxl xlrd).
"""

import re
import zipfile
from pathlib import Path


def clean(value):
    """Cell value as text with collapsed whitespace; '' for an empty cell."""
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # xlrd returns every number as a float
    return re.sub(r"\s+", " ", str(value)).strip() if value is not None else ""


class Sheet:
    def __init__(self, name, cells, merges):
        self.name = name
        self.cells = cells  # list of rows, each a list of raw values (None for empty cells)
        self.merges = merges  # (first row, last row, first col, last col), 0-based and inclusive
        self.nrows = len(cells)
        self.ncols = max((len(row) for row in cells), default=0)

    def get(self, r, c):
        if 0 <= r < self.nrows and 0 <= c < len(self.cells[r]):
            return self.cells[r][c]
        return None

    def text(self, r, c):
        return clean(self.get(r, c))


def load(path):
    """All worksheets of a workbook, including hidden ones."""
    path = Path(path)
    if zipfile.is_zipfile(path):
        import openpyxl
        workbook = openpyxl.load_workbook(path, data_only=True)
        return [Sheet(ws.title, [list(row) for row in ws.iter_rows(values_only=True)],
                      [(m.min_row - 1, m.max_row - 1, m.min_col - 1, m.max_col - 1) for m in ws.merged_cells.ranges])
                for ws in workbook.worksheets]
    import xlrd
    workbook = xlrd.open_workbook(path, formatting_info=True)
    return [Sheet(ws.name, [[(v if v != "" else None) for v in ws.row_values(i)] for i in range(ws.nrows)],
                  [(r0, r1 - 1, c0, c1 - 1) for r0, r1, c0, c1 in ws.merged_cells])
            for ws in workbook.sheets()]
