#!/usr/bin/env python3
"""Download the tourism source data and write the CSV files in data/.

Usage:
    python3 scripts/fetch_tourism_data.py

Sources:
    ELSTAT  "Έρευνα Δραστηριότητας Τουριστικών Καταλυμάτων" (STO12), annual 2025 tables:
            https://www.statistics.gr/el/statistics/-/publication/STO12/2025
    Eurostat tour_occ_nin2 (annual) and tour_occ_nin2m (monthly), nights spent at tourist
            accommodation establishments by NUTS 2 region.

The ELSTAT spreadsheets are kept under data/elstat_tourism_2025/ so the CSVs can be rebuilt
without the website. Needs openpyxl (pip install openpyxl). The dashboard data is then embedded
with scripts/build_tourism_data.py.
"""

import csv
import html
import json
import re
import subprocess
import sys
import urllib.parse
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
RAW = DATA / "elstat_tourism_2025"

YEAR = 2025
PUBLICATION_URL = f"https://www.statistics.gr/el/statistics/-/publication/STO12/{YEAR}"

# ELSTAT document ids of the 2025 publication tables (see the publication page).
DOCS = {
    "table03_arrivals_by_region.xlsx": "600505",
    "table04_nights_by_region.xlsx": "600504",
    "table07_bed_occupancy_by_month.xlsx": "600511",
    "table08_arrivals_by_month.xlsx": "600512",
    "table09_nights_by_month.xlsx": "600513",
    "table14_arrivals_nights_by_origin.xlsx": "600514",
}

EUROSTAT = "https://ec.europa.eu/eurostat/api/dissemination/statistics/1.0/data"

REGION_NAMES = {
    "EL": "Σύνολο Ελλάδας",
    "EL30": "Αττική",
    "EL41": "Βόρειο Αιγαίο",
    "EL42": "Νότιο Αιγαίο",
    "EL43": "Κρήτη",
    "EL51": "Ανατολική Μακεδονία και Θράκη",
    "EL52": "Κεντρική Μακεδονία",
    "EL53": "Δυτική Μακεδονία",
    "EL54": "Ήπειρος",
    "EL61": "Θεσσαλία",
    "EL62": "Ιόνια Νησιά",
    "EL63": "Δυτική Ελλάδα",
    "EL64": "Στερεά Ελλάδα",
    "EL65": "Πελοπόννησος",
}

# ELSTAT writes region names in capitals, without accents, and not always in the same case.
REGION_CODES = {
    "ΣΥΝΟΛΟ": "EL",
    "ΑΝΑΤΟΛΙΚΗΣ ΜΑΚΕΔΟΝΙΑΣ ΚΑΙ ΘΡΑΚΗΣ": "EL51", "ΑΝΑΤΟΛΙΚΗ ΜΑΚΕΔΟΝΙΑ, ΘΡΑΚΗ": "EL51",
    "ΚΕΝΤΡΙΚΗΣ ΜΑΚΕΔΟΝΙΑΣ": "EL52", "ΚΕΝΤΡΙΚΗ ΜΑΚΕΔΟΝΙΑ": "EL52",
    "ΔΥΤΙΚΗΣ ΜΑΚΕΔΟΝΙΑΣ": "EL53", "ΔΥΤΙΚΗ ΜΑΚΕΔΟΝΙΑ": "EL53",
    "ΗΠΕΙΡΟΥ": "EL54", "ΗΠΕΙΡΟΣ": "EL54",
    "ΘΕΣΣΑΛΙΑΣ": "EL61", "ΘΕΣΣΑΛΙΑ": "EL61",
    "ΙΟΝΙΩΝ ΝΗΣΩΝ": "EL62", "ΙΟΝΙΑ ΝΗΣΙΑ": "EL62",
    "ΔΥΤΙΚΗΣ ΕΛΛΑΔΑΣ": "EL63", "ΔΥΤΙΚΗ ΕΛΛΑΔΑ": "EL63",
    "ΣΤΕΡΕΑΣ ΕΛΛΑΔΑΣ": "EL64", "ΣΤΕΡΕΑ ΕΛΛΑΔΑ": "EL64",
    "ΠΕΛΟΠΟΝΝΗΣΟΥ": "EL65", "ΠΕΛΟΠΟΝΝΗΣΟΣ": "EL65",
    "ΑΤΤΙΚΗΣ": "EL30", "ΑΤΤΙΚΗ": "EL30",
    "ΒΟΡΕΙΟΥ ΑΙΓΑΙΟΥ": "EL41", "ΒΟΡΕΙΟ ΑΙΓΑΙΟ": "EL41",
    "ΝΟΤΙΟΥ ΑΙΓΑΙΟΥ": "EL42", "ΝΟΤΙΟ ΑΙΓΑΙΟ": "EL42",
    "ΚΡΗΤΗΣ": "EL43", "ΚΡΗΤΗ": "EL43",
}

ACCOMMODATION_SHEETS = {"Ξενοδοχεία": "hotels", "Ενοικιαζόμενα": "short_stay", "Κάμπινγκ": "camping"}
MONTHS = ["Ιανουάριος", "Φεβρουάριος", "Μάρτιος", "Απρίλιος", "Μάιος", "Ιούνιος", "Ιούλιος",
          "Αύγουστος", "Σεπτέμβριος", "Οκτώβριος", "Νοέμβριος", "Δεκέμβριος"]


def get(url):
    """Fetch a URL with curl; the ELSTAT site answers plain urllib requests with a redirect to http."""
    result = subprocess.run(["curl", "-sSL", "--fail", "--max-time", "120", url], capture_output=True)
    if result.returncode != 0:
        sys.exit(f"download failed ({result.returncode}): {url}\n{result.stderr.decode('utf-8', 'ignore')}")
    return result.stdout


def clean(text):
    return re.sub(r"\s+", " ", str(text)).strip()


def region_code(name):
    key = clean(re.sub(r"\(.*?\)", "", str(name))).upper()
    if key not in REGION_CODES:
        raise KeyError(f"unknown region name: {name!r}")
    return REGION_CODES[key]


def download_elstat():
    RAW.mkdir(parents=True, exist_ok=True)
    page = get(PUBLICATION_URL).decode("utf-8", "ignore")
    links = {}
    for match in re.finditer(r'href="([^"]*documentID=(\d+)[^"]*)"', page):
        links[match.group(2)] = html.unescape(match.group(1))
    for filename, doc_id in DOCS.items():
        target = RAW / filename
        if target.exists():
            continue
        if doc_id not in links:
            sys.exit(f"ELSTAT document {doc_id} ({filename}) not found on {PUBLICATION_URL}")
        target.write_bytes(get(links[doc_id]))
        print("downloaded", filename)


def numbers(row, start, count):
    return [row[i] if isinstance(row[i], (int, float)) else None for i in range(start, start + count)]


def parse_regions():
    """Hotels, short-stay rentals and camping by region and regional unit (tables 3 and 4)."""
    records = {}
    for filename, metric in (("table03_arrivals_by_region.xlsx", "arrivals"),
                             ("table04_nights_by_region.xlsx", "nights")):
        workbook = openpyxl.load_workbook(RAW / filename, data_only=True)
        for sheet_name, accommodation in ACCOMMODATION_SHEETS.items():
            sheet = workbook[sheet_name]
            for row in sheet.iter_rows(min_row=6, values_only=True):
                if not row[0] or not isinstance(row[2], (int, float)):
                    continue
                region_raw, unit_raw = clean(row[0]), clean(row[1] or "")
                code = region_code(region_raw)
                if code == "EL":
                    level, unit = "country", ""
                elif unit_raw.upper() == "ΣΥΝΟΛΟ":
                    level, unit = "region", ""
                else:
                    level, unit = "unit", unit_raw
                key = (accommodation, code, unit)
                entry = records.setdefault(key, {"level": level, "beds": None})
                entry[f"{metric}_domestic"], entry[f"{metric}_foreign"], entry[f"{metric}_total"] = numbers(row, 2, 3)
                if metric == "arrivals" and accommodation == "hotels" and isinstance(row[5], (int, float)):
                    entry["beds"] = row[5]
    path = DATA / "tourism_elstat_2025_regions.csv"
    fields = ["year", "accommodation", "nuts2", "region", "unit", "level", "arrivals_domestic", "arrivals_foreign",
              "arrivals_total", "nights_domestic", "nights_foreign", "nights_total", "beds"]
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, lineterminator="\n")
        writer.writeheader()
        for (accommodation, code, unit), entry in sorted(records.items(), key=lambda kv: (kv[0][0], kv[0][1], kv[0][2])):
            writer.writerow({"year": YEAR, "accommodation": accommodation, "nuts2": code, "region": REGION_NAMES[code],
                             "unit": unit, **{k: ("" if v is None else v) for k, v in entry.items()}})
    print("wrote", path.name, len(records), "rows")


def parse_monthly():
    """Arrivals and nights by month, accommodation type and residence (tables 8 and 9)."""
    rows = {}
    for filename, metric in (("table08_arrivals_by_month.xlsx", "arrivals"), ("table09_nights_by_month.xlsx", "nights")):
        sheet = openpyxl.load_workbook(RAW / filename, data_only=True).worksheets[0]
        for row in sheet.iter_rows(min_row=1, values_only=True):
            if not row[0] or clean(row[0]) not in MONTHS:
                continue
            month = MONTHS.index(clean(row[0])) + 1
            groups = {"domestic": row[1:5], "foreign": row[5:9], "total": row[9:13]}
            for residence, values in groups.items():
                for accommodation, value in zip(("hotels", "camping", "short_stay", "all"), values):
                    rows.setdefault((month, accommodation, residence), {})[metric] = value
    path = DATA / "tourism_elstat_2025_monthly.csv"
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(["year", "month", "accommodation", "residence", "arrivals", "nights"])
        for (month, accommodation, residence), values in sorted(rows.items()):
            writer.writerow([YEAR, month, accommodation, residence, values.get("arrivals", ""), values.get("nights", "")])
    print("wrote", path.name, len(rows), "rows")

    path = DATA / "tourism_elstat_2025_occupancy.csv"
    sheet = openpyxl.load_workbook(RAW / "table07_bed_occupancy_by_month.xlsx", data_only=True).worksheets[0]
    count = 0
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(["year", "month", "bed_occupancy_pct"])
        for row in sheet.iter_rows(values_only=True):
            if row[0] and clean(row[0]) in MONTHS:
                writer.writerow([YEAR, MONTHS.index(clean(row[0])) + 1, row[1]])
                count += 1
    print("wrote", path.name, count, "rows")


def parse_origin():
    """Arrivals and nights by region of destination and country of origin (table 14)."""
    sheet = openpyxl.load_workbook(RAW / "table14_arrivals_nights_by_origin.xlsx", data_only=True).worksheets[0]
    path = DATA / "tourism_elstat_2025_origin.csv"
    count = 0
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(["year", "nuts2", "region", "origin", "indent", "arrivals", "nights"])
        for row in sheet.iter_rows(min_row=4, values_only=True):
            if not row[0] or not isinstance(row[2], (int, float)):
                continue
            region_raw = clean(row[0])
            if region_raw.upper() == "ΠΕΡΙΦΕΡΕΙΑ (NUTS 2)":
                continue
            code = "EL" if region_raw == "Σύνολο" else region_code(region_raw)
            origin = str(row[1])
            indent = (len(origin) - len(origin.lstrip())) // 2
            writer.writerow([YEAR, code, REGION_NAMES[code], clean(origin), indent, row[2], row[3]])
            count += 1
    print("wrote", path.name, count, "rows")


def eurostat(dataset, **filters):
    params = [("format", "JSON"), ("lang", "EN")]
    for key, values in filters.items():
        params += [(key, value) for value in (values if isinstance(values, list) else [values])]
    data = json.loads(get(f"{EUROSTAT}/{dataset}?{urllib.parse.urlencode(params)}"))
    ids, sizes = data["id"], data["size"]
    positions = {dim: list(data["dimension"][dim]["category"]["index"]) for dim in ids}
    for flat, value in data["value"].items():
        flat, record = int(flat), {}
        for dim, size in reversed(list(zip(ids, sizes))):
            flat, index = divmod(flat, size)
            record[dim] = positions[dim][index]
        yield record, value
    return data.get("updated")


def fetch_eurostat():
    geos = list(REGION_NAMES)
    common = dict(geo=geos, nace_r2="I551-I553", unit="NR", c_resid=["TOTAL", "DOM", "FOR"])
    residence = {"TOTAL": "total", "DOM": "domestic", "FOR": "foreign"}

    rows = sorted((r["time"], r["geo"], residence[r["c_resid"]], int(v)) for r, v in eurostat("tour_occ_nin2", **common))
    path = DATA / "tourism_eurostat_nights_annual.csv"
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(["year", "nuts2", "region", "residence", "nights"])
        for year, geo, res, value in rows:
            writer.writerow([year, geo, REGION_NAMES[geo], res, value])
    print("wrote", path.name, len(rows), "rows")

    rows = sorted((r["time"], r["month"][1:], r["geo"], residence[r["c_resid"]], int(v))
                  for r, v in eurostat("tour_occ_nin2m", **common, month=[f"M{m:02d}" for m in range(1, 13)]))
    path = DATA / "tourism_eurostat_nights_monthly.csv"
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(["year", "month", "nuts2", "region", "residence", "nights"])
        for year, month, geo, res, value in rows:
            writer.writerow([year, int(month), geo, REGION_NAMES[geo], res, value])
    print("wrote", path.name, len(rows), "rows")


if __name__ == "__main__":
    download_elstat()
    parse_regions()
    parse_monthly()
    parse_origin()
    fetch_eurostat()
