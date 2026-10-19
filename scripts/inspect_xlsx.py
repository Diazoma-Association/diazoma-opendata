#!/usr/bin/env python3
"""Print the structure of an XLSX file without loading it into memory: sheets, sizes, header rows and sample rows.

Usage:
    python3 scripts/inspect_xlsx.py FILE.xlsx [--rows N] [--sheet NAME]

Use it on a downloaded file (for example the ESPA list of operations) before writing or adjusting an importer:
the column names must come from the real file, never from assumptions. The file is read in streaming mode, so
large workbooks do not exhaust the memory.
"""

import argparse
import itertools

import openpyxl


def main():
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("file")
    parser.add_argument("--rows", type=int, default=8, help="number of rows to show per sheet (default 8)")
    parser.add_argument("--sheet", help="only this sheet")
    args = parser.parse_args()

    workbook = openpyxl.load_workbook(args.file, read_only=True, data_only=True)
    for sheet in workbook.worksheets:
        if args.sheet and sheet.title != args.sheet:
            continue
        print(f"== sheet {sheet.title!r}: {sheet.max_row} rows x {sheet.max_column} columns")
        for number, row in enumerate(itertools.islice(sheet.iter_rows(values_only=True), args.rows), 1):
            cells = [str(value)[:40] if value is not None else "" for value in row]
            while cells and not cells[-1]:
                cells.pop()
            print(f"{number:>4}: {cells}")
    workbook.close()


if __name__ == "__main__":
    main()
